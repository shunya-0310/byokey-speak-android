from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT.parent / "feature-graphic-2026-07-30" / "assets"
OUT = ROOT / "byokey-speak-play-screenshots-editable-light-5pages.pptx"
LOGO_KEY = ROOT / "assets-keyhole-transparent.png"

W, H = 1080, 1920
SCALE = 100.0

PARCHMENT = RGBColor(242, 232, 208)
PARCHMENT_2 = RGBColor(232, 216, 185)
INK = RGBColor(48, 39, 31)
INK_SOFT = RGBColor(95, 79, 61)
GOLD = RGBColor(165, 124, 45)
GOLD_DARK = RGBColor(116, 84, 31)
CORAL = RGBColor(177, 92, 77)
MOSS = RGBColor(91, 110, 81)
PHONE = RGBColor(250, 244, 228)
PHONE_SCREEN = RGBColor(255, 251, 240)
CARD = RGBColor(247, 238, 218)

FONT_HEAD = "Noto Sans JP"
FONT_BODY = "Noto Sans JP"
FONT_SERIF = "Noto Serif JP"


def u(px):
    return Inches(px / SCALE)


def make_keyhole_asset():
    logo = Image.open(ASSETS / "byokey-lab-logo.png").convert("RGBA")
    # Crop to the keyhole mark in the source logo.
    crop = logo.crop((392, 320, 635, 665))
    data = crop.load()
    for y in range(crop.height):
        for x in range(crop.width):
            r, g, b, a = data[x, y]
            if r > 246 and g > 246 and b > 246:
                data[x, y] = (255, 255, 255, 0)
    crop.save(LOGO_KEY)


def add_text(slide, x, y, w, h, text, size, color=INK, bold=True, font=FONT_BODY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    box.text_frame.clear()
    box.text_frame.margin_left = 0
    box.text_frame.margin_right = 0
    box.text_frame.margin_top = 0
    box.text_frame.margin_bottom = 0
    box.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = box.text_frame.paragraphs[0] if idx == 0 else box.text_frame.add_paragraph()
        p.alignment = align
        p.text = ""
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=True, line_width=1.5, transparency=0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        u(x),
        u(y),
        u(w),
        u(h),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_width)
    else:
        shp.line.fill.background()
    return shp


def add_line(slide, x1, y1, x2, y2, color=GOLD, width=3):
    line = slide.shapes.add_connector(1, u(x1), u(y1), u(x2), u(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_background(slide):
    add_rect(slide, 0, 0, W, H, PARCHMENT, radius=False)
    # Editable, subtle parchment-like bands. Kept light so text stays readable.
    for y in range(70, H, 120):
        add_line(slide, 0, y, W, y + 8, RGBColor(217, 198, 158), 0.6)
    for x, y, s, c in [
        (905, 110, 7, GOLD),
        (135, 250, 5, GOLD),
        (970, 440, 5, MOSS),
        (90, 1540, 4, CORAL),
        (810, 1700, 5, GOLD),
    ]:
        add_text(slide, x, y, 30, 30, "✦", 14 + s, c, True, FONT_SERIF)


def add_brand(slide):
    slide.shapes.add_picture(str(LOGO_KEY), u(66), u(60), height=u(112))
    add_text(slide, 190, 68, 500, 50, "BYOKey Speak", 33, INK, True, FONT_SERIF)
    add_text(slide, 192, 118, 490, 34, "The key is in your hand.", 19, GOLD_DARK, True, FONT_SERIF)


def add_phone(slide, label, items, accent_item=None):
    x, y, w, h = 242, 740, 596, 1375
    # Bottom intentionally extends past the 1920px slide by about 10%.
    add_rect(slide, x + 18, y + 24, w, h, RGBColor(186, 169, 132), None, True, transparency=42)
    add_rect(slide, x, y, w, h, RGBColor(44, 39, 35), GOLD_DARK, True, 3)
    add_rect(slide, x + 24, y + 24, w - 48, h - 48, PHONE, GOLD, True, 1.3)
    add_rect(slide, x + 206, y + 22, 184, 18, RGBColor(64, 57, 51), None, True)
    add_rect(slide, x + 50, y + 70, w - 100, h - 150, PHONE_SCREEN, RGBColor(211, 181, 116), True, 1.2)
    add_text(slide, x + 82, y + 108, w - 164, 48, label, 25, GOLD_DARK, True, FONT_BODY)
    add_line(slide, x + 82, y + 170, x + w - 82, y + 170, GOLD, 1.6)
    cy = y + 216
    for idx, (title, body) in enumerate(items):
        box_h = 128 if len(body) < 28 else 154
        add_rect(slide, x + 82, cy, w - 164, box_h, CARD, GOLD, True, 1.25)
        title_color = CORAL if idx == accent_item else GOLD_DARK
        add_text(slide, x + 108, cy + 22, w - 216, 32, title, 18, title_color, True, FONT_BODY)
        add_text(slide, x + 108, cy + 58, w - 216, box_h - 66, body, 20, INK, True, FONT_BODY)
        cy += box_h + 28
    add_rect(slide, x + 96, y + 1118, w - 192, 58, RGBColor(238, 225, 193), GOLD, True, 1)
    add_text(slide, x + 128, y + 1133, w - 256, 28, "ここに実画面スクショを差し替え", 16, INK_SOFT, True, FONT_BODY, PP_ALIGN.CENTER)


def add_badge(slide, text, x, y, w=590):
    add_rect(slide, x, y, w, 58, RGBColor(255, 249, 232), GOLD_DARK, True, 1.6)
    add_text(slide, x + 24, y + 13, w - 48, 32, text, 20, INK, True, FONT_BODY, PP_ALIGN.CENTER)


slides = [
    {
        "lines": [("英語で話す。", INK), ("言葉に詰まっても、", INK), ("日本語で聞ける。", CORAL)],
        "sub": "先生に相談するように、そのまま会話を続けられます。",
        "phone": "通常チャット画面",
        "items": [("英語の会話", "I missed the train."), ("日本語で相談", "“乗り遅れた”って自然に言うと？"), ("コーチの返答", "英語表現＋日本語の短い説明")],
        "badge": "アプリは買い切り｜AI利用料は使った分だけ",
    },
    {
        "lines": [("話し方も、教え方も。", INK), ("あなたの言葉で決められる。", CORAL)],
        "sub": "性格・添削方針・説明言語・会話スタイルを自由に設定",
        "phone": "Coach Personalities & Skills",
        "items": [("設定欄", "間違いは日本語で短く説明してください。"), ("添削方針", "会話を止めず、最後にまとめて添削。"), ("話し方", "質問を多めにして、やさしく励ます。")],
    },
    {
        "lines": [("返信を選ぶのではなく、", INK), ("言いたいことを英語にする。", CORAL)],
        "sub": "日本語でも、英語でも、混ざっていても大丈夫です。",
        "phone": "Quick Assist",
        "items": [("入力", "子どもを寝かしつけないといけない"), ("候補 1", "I need to put my kids to bed."), ("候補 2", "It’s time for bed, so I have to go.")],
        "accent_item": 0,
    },
    {
        "lines": [("話したあとに、", INK), ("次に伸ばすところが", INK), ("見えてくる。", CORAL)],
        "sub": "会話履歴から、強み・傾向・改善ポイントを分析",
        "phone": "会話分析結果",
        "items": [("強み", "理由を添えて話せている"), ("よくある傾向", "表現が直訳になりやすい"), ("改善ポイント", "短い自然な言い換えを増やす"), ("次の練習項目", "反対意見をやさしく伝える")],
    },
    {
        "lines": [("使うAIも、", INK), ("毎月の費用も、", INK), ("自分で選ぶ。", CORAL)],
        "sub": "アプリは買い切り。AI利用料は使った分だけ。",
        "phone": "AIサービス・モデル選択画面",
        "items": [("LLM provider", "Gemini / OpenAI / Claude"), ("Model", "会話品質や費用感で選択"), ("API key", "対応AIサービスのAPIキーを入力")],
        "note": "利用には対応AIサービスのAPIキーが必要です。",
        "small": "会話履歴は原則端末内に保存",
    },
]


make_keyhole_asset()

prs = Presentation()
prs.slide_width = u(W)
prs.slide_height = u(H)
blank = prs.slide_layouts[6]

for idx, spec in enumerate(slides, start=1):
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_brand(slide)
    add_text(slide, 72, 198, 120, 42, f"{idx}/5", 23, GOLD_DARK, True, FONT_SERIF)
    y = 276
    for line, color in spec["lines"]:
        add_text(slide, 72, y, 920, 78, line, 45, color, True, FONT_HEAD)
        y += 78
    add_line(slide, 72, y + 8, 520, y + 8, GOLD_DARK, 4)
    add_text(slide, 72, y + 78, 780, 78, spec["sub"], 24, INK_SOFT, True, FONT_BODY)
    add_phone(slide, spec["phone"], spec["items"], spec.get("accent_item"))
    if spec.get("badge"):
        add_badge(slide, spec["badge"], 72, 1768, 620)
    if spec.get("note"):
        add_text(slide, 72, 1728, 820, 34, spec["note"], 18, INK_SOFT, True, FONT_BODY)
    if spec.get("small"):
        add_badge(slide, spec["small"], 72, 1810, 430)

prs.save(OUT)
print(OUT)
