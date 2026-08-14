from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE = Path("byokey-speak-android/docs/marketing/play-screenshots-wireframes-2026-07-31")
ATTACH = Path(
    ".codex-remote-attachments/"
    "019fb222-aba3-7352-8189-59f0acb45451/"
    "76e4f069-b1eb-45d6-8d7c-57d452ae8706/"
    "1-Photo-1.jpg"
)

OUT = BASE / "byokey-speak-feature-split-2pages-1080x1920-editable.pptx"
PREVIEW = BASE / "byokey-speak-feature-split-2pages-2160x1920-preview.png"
BG = BASE / "feature-panorama-bg.png"
KEYHOLE = BASE / "assets-keyhole-transparent.png"

W, H = 1080, 1920
SOURCE_W, SOURCE_H = 2073.6, 1843.2
SCALE = H / SOURCE_H
SPLIT_X = SOURCE_W / 2

NAVY = RGBColor(8, 20, 35)
CARD = RGBColor(29, 50, 78)
GOLD = RGBColor(185, 145, 55)
GOLD_SOFT = RGBColor(219, 193, 111)
CORAL = RGBColor(232, 151, 132)
WHITE = RGBColor(248, 245, 235)
MUTED = RGBColor(210, 194, 150)

FONT_HEAD = "Noto Sans JP"
FONT_SERIF = "Noto Serif JP"


def u(px: float):
    return Inches(px / 100)


def sx(x: float, offset: float) -> float:
    return (x - offset) * SCALE


def sy(y: float) -> float:
    return y * SCALE


def sw(w: float) -> float:
    return w * SCALE


def add_text(slide, x, y, w, h, text, size, color=WHITE, bold=True, font=FONT_HEAD, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=True, line_width=1.0, transparency=0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        u(x),
        u(y),
        u(w),
        u(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    return shape


def add_line(slide, x1, y1, x2, y2, color=GOLD, width=3):
    line = slide.shapes.add_connector(1, u(x1), u(y1), u(x2), u(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


TEXTS = [
    {
        "text": "BYOKey Speak",
        "x": 87.34569952058348,
        "y": 108.88,
        "w": 420.375,
        "h": 68,
        "size": 43,
        "color": WHITE,
        "font": FONT_SERIF,
    },
    {
        "text": "AI英会話は、“Bring Your Own Key”で新たな体験へ",
        "x": 88.32,
        "y": 184.2,
        "w": 699.828125,
        "h": 35,
        "size": 21,
        "color": GOLD_SOFT,
        "font": FONT_HEAD,
    },
    {
        "text": "英語で話す。",
        "x": 88.32,
        "y": 447.346592,
        "w": 864,
        "h": 113.28,
        "size": 68,
        "color": WHITE,
        "font": FONT_HEAD,
    },
    {
        "text": "言葉に詰まっても、",
        "x": 88.32,
        "y": 639.826592,
        "w": 815.953125,
        "h": 108,
        "size": 68,
        "color": WHITE,
        "font": FONT_HEAD,
    },
    {
        "text": "日本語で聞ける。",
        "x": 88.32,
        "y": 827.026592,
        "w": 816,
        "h": 113.28,
        "size": 72,
        "color": CORAL,
        "font": FONT_HEAD,
    },
    {
        "text": "話し方も、教え方も、自分好みに。",
        "x": 88.32,
        "y": 1055.966592,
        "w": 725.28125,
        "h": 54,
        "size": 29,
        "color": MUTED,
        "font": FONT_HEAD,
    },
]

BADGE = {
    "text": "アプリは買い切り、\n毎月の費用は使った分だけ",
    "x": 88.32,
    "y": 1225.966592,
    "w": 816,
    "h": 168.346816,
}

KEYHOLE_POS = {
    "x": 956.2024039684652,
    "y": 717.4429193962378,
    "w": 287.5951918070387,
    "h": 408.3141612075241,
}

PHONE_POS = {
    "x": 1159.8199204813868,
    "y": -20.2,
    "w": 842.3200397593066,
    "h": 1894.16,
}


def build_pptx():
    prs = Presentation()
    prs.slide_width = u(W)
    prs.slide_height = u(H)

    for offset, label in [(0, "01_left"), (SPLIT_X, "02_right")]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # The background is one continuous 2160x1920 image. Each page exposes one half.
        slide.shapes.add_picture(str(BG), u(-offset * SCALE), u(0), width=u(2160), height=u(1920))

        # Subtle accent strokes. They are editable vector lines and will clip naturally.
        for x, y, w, color in [
            (1240, 365, 66, GOLD),
            (1325, 505, 88, CORAL),
            (1260, 640, 72, RGBColor(161, 177, 119)),
            (1820, 166, 132, GOLD),
        ]:
            add_line(slide, x - (1080 if offset else 0), y, x + w - (1080 if offset else 0), y, color, 2)

        # Copy block and badge. Keeping them on both slides with true x-offset means
        # the right page clips them out, but the elements remain editable if selected.
        for item in TEXTS:
            add_text(
                slide,
                sx(item["x"], offset),
                sy(item["y"]),
                sw(item["w"]),
                sw(item["h"]),
                item["text"],
                item["size"],
                item["color"],
                True,
                item["font"],
            )

        # Underline below the emphasized phrase.
        add_line(
            slide,
            sx(88.32, offset),
            sy(970),
            sx(770, offset),
            sy(970),
            GOLD,
            5,
        )

        badge_x, badge_y = sx(BADGE["x"], offset), sy(BADGE["y"] - 18)
        badge_w, badge_h = sw(BADGE["w"]), sw(BADGE["h"] + 36)
        add_rect(slide, badge_x, badge_y, badge_w, badge_h, CARD, GOLD, True, 1.5)
        add_text(
            slide,
            sx(BADGE["x"] + 24, offset),
            sy(BADGE["y"] + 24),
            sw(BADGE["w"] - 48),
            sw(BADGE["h"] - 24),
            BADGE["text"],
            31,
            WHITE,
            True,
            FONT_HEAD,
            PP_ALIGN.CENTER,
        )

        # Central keyhole. It intentionally straddles the split line.
        slide.shapes.add_picture(
            str(KEYHOLE),
            u(sx(KEYHOLE_POS["x"], offset)),
            u(sy(KEYHOLE_POS["y"])),
            width=u(sw(KEYHOLE_POS["w"])),
            height=u(sw(KEYHOLE_POS["h"])),
        )

        # Actual app screenshot, positioned from the user's current Canva design.
        slide.shapes.add_picture(
            str(ATTACH),
            u(sx(PHONE_POS["x"], offset)),
            u(sy(PHONE_POS["y"])),
            width=u(sw(PHONE_POS["w"])),
            height=u(sw(PHONE_POS["h"])),
        )

        slide.name = label

    prs.save(OUT)


def build_preview():
    canvas = Image.open(BG).convert("RGB")
    key = Image.open(KEYHOLE).convert("RGBA")
    phone = Image.open(ATTACH).convert("RGB")

    def place(img, box):
        x, y, w, h = box
        resized = img.resize((round(w), round(h)), Image.Resampling.LANCZOS)
        if resized.mode == "RGBA":
            canvas.alpha_composite(resized, (round(x), round(y)))
        else:
            canvas.paste(resized, (round(x), round(y)))

    canvas = canvas.convert("RGBA")
    place(
        key,
        (
            sy(KEYHOLE_POS["x"]),
            sy(KEYHOLE_POS["y"]),
            sw(KEYHOLE_POS["w"]),
            sw(KEYHOLE_POS["h"]),
        ),
    )
    place(
        phone,
        (
            sy(PHONE_POS["x"]),
            sy(PHONE_POS["y"]),
            sw(PHONE_POS["w"]),
            sw(PHONE_POS["h"]),
        ),
    )
    canvas.convert("RGB").save(PREVIEW)


def main():
    if not BG.exists():
        raise FileNotFoundError(BG)
    if not KEYHOLE.exists():
        raise FileNotFoundError(KEYHOLE)
    if not ATTACH.exists():
        raise FileNotFoundError(ATTACH)
    build_pptx()
    build_preview()
    print(OUT)
    print(PREVIEW)


if __name__ == "__main__":
    main()
